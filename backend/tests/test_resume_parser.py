"""이력서 파일 파서와 multipart API 회귀 테스트."""

from __future__ import annotations

import io
import struct
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services.resume_parser import (
    ResumeParseError,
    _extract_hwp_paragraphs,
    parse_resume,
)

client = TestClient(app)


def make_docx() -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_heading("김웨이 이력서", level=1)
    document.add_paragraph("데이터 분석가로 5년간 고객 행동 분석과 대시보드 자동화를 담당했습니다.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "기술"
    table.cell(0, 1).text = "Python, SQL, Tableau"
    document.save(stream)
    return stream.getvalue()


def make_hwpx() -> bytes:
    stream = io.BytesIO()
    section = """<?xml version="1.0" encoding="UTF-8"?>
    <hs:sec xmlns:hs="urn:hancom:section" xmlns:hp="urn:hancom:paragraph">
      <hp:p><hp:run><hp:t>서비스 기획자로 6년간 사용자 조사와 요구사항 정의를 담당했습니다.</hp:t></hp:run></hp:p>
      <hp:p><hp:run><hp:t>애자일 프로젝트 관리와 데이터 기반 의사결정 경험이 있습니다.</hp:t></hp:run></hp:p>
    </hs:sec>"""
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("mimetype", "application/hwp+zip")
        archive.writestr("Contents/section0.xml", section)
    return stream.getvalue()


def make_text_pdf() -> bytes:
    content = b"BT /F1 12 Tf 72 720 Td (Senior data analyst with five years experience) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode())
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(pdf)


def test_docx_parser_includes_paragraphs_and_tables():
    text, file_type, warnings = parse_resume("resume.docx", make_docx())
    assert file_type == "docx"
    assert "데이터 분석가" in text
    assert "Python, SQL, Tableau" in text
    assert warnings == []


def test_pdf_parser_extracts_text():
    text, file_type, warnings = parse_resume("resume.pdf", make_text_pdf())
    assert file_type == "pdf"
    assert "Senior data analyst" in text
    assert warnings == []


def test_hwpx_parser_extracts_paragraphs():
    text, file_type, warnings = parse_resume("resume.hwpx", make_hwpx())
    assert file_type == "hwpx"
    assert "서비스 기획자" in text
    assert "데이터 기반 의사결정" in text
    assert warnings == []


def test_hwp_paragraph_record_decoder():
    text = "고객 상담 업무와 민원 해결을 7년간 담당했습니다."
    payload = text.encode("utf-16le")
    record_header = 67 | (len(payload) << 20)
    assert _extract_hwp_paragraphs(struct.pack("<I", record_header) + payload) == [text]


@pytest.mark.parametrize("filename", ["resume.txt", "resume.exe", "resume"])
def test_unsupported_extension_is_rejected(filename: str):
    with pytest.raises(ResumeParseError, match="PDF, DOCX"):
        parse_resume(filename, b"plain text resume contents long enough")


def test_spoofed_pdf_is_rejected():
    with pytest.raises(ResumeParseError, match="실제 PDF"):
        parse_resume("resume.pdf", b"not a pdf file but long enough")


def test_parse_file_endpoint():
    response = client.post(
        "/api/diagnosis/parse-file",
        files={
            "file": (
                "career.docx",
                make_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    body = response.json()
    assert body["filename"] == "career.docx"
    assert body["file_type"] == "docx"
    assert body["char_count"] == len(body["text"])
    assert "데이터 분석가" in body["text"]


def test_parse_file_endpoint_rejects_invalid_file():
    response = client.post(
        "/api/diagnosis/parse-file",
        files={"file": ("resume.pdf", b"fake pdf content", "application/pdf")},
    )
    assert response.status_code == 422
    assert "실제 PDF" in response.json()["detail"]
