"""이력서 파일에서 진단에 사용할 일반 텍스트를 안전하게 추출한다."""

from __future__ import annotations

import io
import re
import struct
import zlib
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import olefile
from docx import Document
from pypdf import PdfReader

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_UNCOMPRESSED_SIZE = 50 * 1024 * 1024
MAX_TEXT_LENGTH = 50_000
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".hwp", ".hwpx"}
OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class ResumeParseError(ValueError):
    """사용자가 수정할 수 있는 파일 형식/내용 오류."""


def parse_resume(filename: str, data: bytes) -> tuple[str, str, list[str]]:
    """파일명과 바이트를 검증한 뒤 (텍스트, 형식, 경고)를 반환한다."""
    safe_name = Path(filename or "").name
    extension = Path(safe_name).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeParseError("PDF, DOCX, HWP, HWPX 파일만 업로드할 수 있습니다.")
    if not data:
        raise ResumeParseError("파일 내용이 비어 있습니다.")
    if len(data) > MAX_FILE_SIZE:
        raise ResumeParseError("파일은 10MB 이하만 업로드할 수 있습니다.")

    warnings: list[str] = []
    if extension == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ResumeParseError("확장자와 실제 PDF 파일 형식이 일치하지 않습니다.")
        text = _parse_pdf(data, warnings)
    elif extension == ".docx":
        _validate_zip(data)
        text = _parse_docx(data)
    elif extension == ".hwpx":
        _validate_zip(data)
        text = _parse_hwpx(data)
    else:
        if not data.startswith(OLE_SIGNATURE):
            raise ResumeParseError("확장자와 실제 HWP 파일 형식이 일치하지 않습니다.")
        text = _parse_hwp(data, warnings)

    text = _normalize_text(text)
    if len(text) < 20:
        raise ResumeParseError(
            "분석할 수 있는 텍스트가 충분하지 않습니다. 스캔 이미지 문서라면 "
            "텍스트가 포함된 PDF·DOCX·HWP 파일을 사용해 주세요."
        )
    if len(text) > MAX_TEXT_LENGTH:
        text = text[:MAX_TEXT_LENGTH].rstrip()
        warnings.append("진단 입력 한도에 맞춰 추출 텍스트를 50,000자로 잘랐습니다.")
    return text, extension.lstrip("."), warnings


def _validate_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if not infos:
                raise ResumeParseError("압축 문서에 읽을 수 있는 내용이 없습니다.")
            if sum(item.file_size for item in infos) > MAX_UNCOMPRESSED_SIZE:
                raise ResumeParseError("압축을 푼 문서 크기가 허용 범위를 초과합니다.")
            for item in infos:
                normalized = item.filename.replace("\\", "/")
                if normalized.startswith("/") or ".." in normalized.split("/"):
                    raise ResumeParseError("안전하지 않은 문서 내부 경로가 포함되어 있습니다.")
    except (zipfile.BadZipFile, RuntimeError) as exc:
        raise ResumeParseError("손상되었거나 암호화된 문서 파일입니다.") from exc


def _parse_pdf(data: bytes, warnings: list[str]) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ResumeParseError("암호가 설정된 PDF는 분석할 수 없습니다.")
        pages: list[str] = []
        empty_pages = 0
        for page in reader.pages:
            page_text = page.extract_text() or ""
            if not page_text.strip():
                empty_pages += 1
            pages.append(page_text)
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError("PDF 텍스트를 읽지 못했습니다.") from exc

    if empty_pages:
        warnings.append(f"텍스트가 감지되지 않은 PDF 페이지가 {empty_pages}쪽 있습니다.")
    return "\n\n".join(pages)


def _parse_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeParseError("DOCX 문서를 읽지 못했습니다.") from exc

    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                blocks.append(" | ".join(values))
    return "\n".join(blocks)


def _parse_hwpx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            if not any(name.lower() == "mimetype" for name in names):
                raise ResumeParseError("올바른 HWPX 문서 구조가 아닙니다.")
            section_names = sorted(
                name
                for name in names
                if re.fullmatch(r"Contents/section\d+\.xml", name, re.IGNORECASE)
            )
            if not section_names:
                raise ResumeParseError("HWPX 본문 섹션을 찾지 못했습니다.")

            paragraphs: list[str] = []
            for name in section_names:
                root = ElementTree.fromstring(archive.read(name))
                for paragraph in root.iter():
                    if _local_name(paragraph.tag) != "p":
                        continue
                    text = "".join(
                        node.text or ""
                        for node in paragraph.iter()
                        if _local_name(node.tag) == "t"
                    ).strip()
                    if text:
                        paragraphs.append(text)
            return "\n".join(paragraphs)
    except ResumeParseError:
        raise
    except (zipfile.BadZipFile, ElementTree.ParseError, KeyError) as exc:
        raise ResumeParseError("HWPX 문서를 읽지 못했습니다.") from exc


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _parse_hwp(data: bytes, warnings: list[str]) -> str:
    try:
        with olefile.OleFileIO(io.BytesIO(data)) as hwp:
            if not hwp.exists("FileHeader"):
                raise ResumeParseError("올바른 HWP 5.x 문서가 아닙니다.")
            header = hwp.openstream("FileHeader").read()
            if not header.startswith(b"HWP Document File"):
                raise ResumeParseError("지원하지 않는 HWP 문서 버전입니다.")
            flags = int.from_bytes(header[36:40], "little")
            if flags & 0x02:
                raise ResumeParseError("암호가 설정된 HWP는 분석할 수 없습니다.")
            compressed = bool(flags & 0x01)

            sections = sorted(
                (
                    path
                    for path in hwp.listdir(streams=True, storages=False)
                    if len(path) == 2
                    and path[0] == "BodyText"
                    and path[1].startswith("Section")
                ),
                key=lambda path: int(path[1].removeprefix("Section") or 0),
            )
            paragraphs: list[str] = []
            for path in sections:
                section = hwp.openstream(path).read()
                if compressed:
                    section = zlib.decompress(section, -15)
                paragraphs.extend(_extract_hwp_paragraphs(section))

            if paragraphs:
                return "\n".join(paragraphs)
            if hwp.exists("PrvText"):
                warnings.append("HWP 본문 대신 미리보기 텍스트를 사용했습니다.")
                return hwp.openstream("PrvText").read().decode("utf-16le", errors="ignore")
    except ResumeParseError:
        raise
    except (OSError, IOError, struct.error, zlib.error, ValueError) as exc:
        raise ResumeParseError("HWP 문서를 읽지 못했습니다.") from exc
    return ""


def _extract_hwp_paragraphs(section: bytes) -> list[str]:
    paragraphs: list[str] = []
    offset = 0
    while offset + 4 <= len(section):
        header = struct.unpack_from("<I", section, offset)[0]
        offset += 4
        tag_id = header & 0x3FF
        size = (header >> 20) & 0xFFF
        if size == 0xFFF:
            if offset + 4 > len(section):
                break
            size = struct.unpack_from("<I", section, offset)[0]
            offset += 4
        if offset + size > len(section):
            break
        payload = section[offset : offset + size]
        offset += size
        if tag_id == 67:  # HWPTAG_PARA_TEXT
            text = _decode_hwp_text(payload).strip()
            if text:
                paragraphs.append(text)
    return paragraphs


def _decode_hwp_text(payload: bytes) -> str:
    units = [
        int.from_bytes(payload[index : index + 2], "little")
        for index in range(0, len(payload) - 1, 2)
    ]
    output: list[str] = []
    index = 0
    while index < len(units):
        code = units[index]
        if code in {10, 13}:
            output.append("\n")
        elif code in {9, 24, 30, 31}:
            output.append(" ")
        elif code < 32:
            # HWP 인라인 확장 제어문자는 제어 코드 포함 8개의 UTF-16 단위다.
            index += 7 if code in {1, 2, 3, 11, 12, 14, 15, 16, 17, 18, 21, 22, 23} else 0
        elif code != 0xFFFF:
            output.append(chr(code))
        index += 1
    return "".join(output)


def _normalize_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    normalized: list[str] = []
    blank = False
    for line in lines:
        if line:
            normalized.append(line)
            blank = False
        elif normalized and not blank:
            normalized.append("")
            blank = True
    return "\n".join(normalized).strip()
